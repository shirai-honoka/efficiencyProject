#全体集計
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime, timedelta, time
from pathlib import Path
from db import get_engine

def export_worktime_total(zexis_keiyaku_no, pattern_code="PAT001"):
    engine = get_engine()

    # 契約情報取得
    sql_order = """
        SELECT 
            zexis_keiyaku_no,
            keiyaku_n,
            keiyaku_start_bi,
            keiyaku_end_bi,
            torihiki_rn,
            time_width_low,
            time_width_up
        FROM d_keiyaku_seikyu_h
        WHERE zexis_keiyaku_no = %(zexis)s
    """
    sql_tanto = """
        SELECT 
            tanto_cd,
            tanto_n
        FROM d_keiyaku_seikyu_m
        WHERE zexis_keiyaku_no = %(zexis)s  
    """
    sql_rcv = """
        SELECT *
        FROM weprj.rcv_params
        WHERE order_no = %(zexis)s AND del_flg = FALSE
    """

    df_order = pd.read_sql(sql_order, con=engine, params={"zexis": zexis_keiyaku_no})
    df_tanto = pd.read_sql(sql_tanto, con=engine, params={"zexis": zexis_keiyaku_no})
    df_rcv = pd.read_sql(sql_rcv, con=engine, params={"zexis": zexis_keiyaku_no})

    if df_order.empty or df_rcv.empty:
        print(f"[ERROR] データ不足: {zexis_keiyaku_no}")
        return

    order_data = df_order.iloc[0]
    rcv_data = df_rcv.iloc[0]

    start_date = pd.to_datetime(order_data["keiyaku_start_bi"])
    end_date = pd.to_datetime(order_data["keiyaku_end_bi"])

    base_dir = Path(__file__).resolve().parent.parent
    template_path = base_dir / "templates" / "全体集計A.xlsx"
    output_dir = base_dir / "output"
    output_dir.mkdir(exist_ok=True)

    wb = load_workbook(template_path)
    ws = wb.active

    # 日付・契約情報
    today = datetime.today()
    ws["A4"]=f"{today.year}年{today.month}月度　実績時間報告書　全体集計"
    ws["E10"] = today.year
    ws["H10"] = today.month
    ws["J10"] = today.day
    ws["D1"] = order_data["torihiki_rn"]
    ws["E7"] = order_data["zexis_keiyaku_no"]
    ws["E8"] = order_data["keiyaku_n"]
    ws["D2"] = ":" + rcv_data["ae_name"]
    ws["Y7"] = "業務担当者：" + rcv_data["pic_name"]
    ws["E9"] = f"{start_date:%Y/%m/%d}～{end_date:%Y/%m/%d}"
    ws["J14"] = order_data["time_width_low"]
    ws["M14"] = order_data["time_width_up"]

    start_row = 20

    # ---------------------------------------
    # personal と同じ秒単位計算
    # ---------------------------------------
    def to_seconds(val):
        if pd.isna(val) or val == "":
            return 0
        if isinstance(val, timedelta):
            return int(val.total_seconds())
        if isinstance(val, time):
            return val.hour*3600 + val.minute*60 + val.second
        if isinstance(val, (int, float)):
            return int(val*3600)  # float時間を秒に
        if isinstance(val, str):
            parts = list(map(int, val.split(":")))
            h = parts[0]
            m = parts[1] if len(parts) > 1 else 0
            s = parts[2] if len(parts) > 2 else 0
            return h*3600 + m*60 + s
        return 0

    # 担当者ごとに集計
    for i, row in df_tanto.iterrows():
        tanto_cd = row["tanto_cd"]
        ws[f"B{start_row + i}"] = row["tanto_n"]
        ws[f"AC{start_row + i}"] = "なし"

        df_work = pd.read_sql(
            "SELECT * FROM weprj.work_record WHERE del_flg = false AND member_no = %s",
            con=engine,
            params=(tanto_cd,),
        )

        if df_work.empty:
            continue

        df_work['work_time'] = df_work['work_time'].apply(to_seconds)
        df_work['work_date'] = pd.to_datetime(df_work['work_date'], format="%Y%m%d", errors='coerce')
        df_work = df_work[(df_work['work_date'] >= start_date) & (df_work['work_date'] <= end_date)]

        # 月間作業時間
        monthly_seconds = df_work['work_time'].sum()
        monthly_hours = monthly_seconds / 3600
        ws[f"K{start_row + i}"] = monthly_hours

        # 超過・控除
        time_low = order_data["time_width_low"]
        time_up = order_data["time_width_up"]
        if monthly_hours < time_low:
            ws[f"N{start_row + i}"] = monthly_hours - time_low
        elif monthly_hours > time_up:
            ws[f"N{start_row + i}"] = monthly_hours - time_up
        else:
            ws[f"N{start_row + i}"] = 0.0

        # 深夜時間
        df_night = df_work.dropna(subset=['night_time'])
        df_night = df_night[df_night['night_time'] != ""].drop_duplicates(subset=['work_date'])
        night_seconds = df_night['night_time'].apply(to_seconds).sum()
        ws[f"Q{start_row + i}"] = len(df_night)
        ws[f"T{start_row + i}"] = night_seconds / 3600

        # 休日時間
        df_holiday = df_work.dropna(subset=['paid_leave_time'])
        df_holiday = df_holiday[df_holiday['paid_leave_time'] != ""].drop_duplicates(subset=['work_date'])
        holiday_seconds = df_holiday['paid_leave_time'].apply(to_seconds).sum()
        ws[f"W{start_row + i}"] = len(df_holiday)
        ws[f"Z{start_row + i}"] = holiday_seconds / 3600

    output_path = output_dir / f"全体集計A_{zexis_keiyaku_no}.xlsx"
    wb.save(output_path)
    print(f"全体集計を出力: {output_path}")
