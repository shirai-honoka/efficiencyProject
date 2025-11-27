#個人別詳細
import pandas as pd
from db import get_engine  # DB接続
from openpyxl import load_workbook
from openpyxl.styles import Font
from datetime import datetime, timedelta, time
from pathlib import Path

def time_to_str(t):
    """時間をHH:MM形式に変換"""
    if pd.isna(t) or t in ("", None):
        return ""
    if isinstance(t, (pd.Timestamp, datetime)):
        return t.strftime('%H:%M')
    if isinstance(t, (pd.Timedelta, timedelta)):
        total_seconds = t.total_seconds()
        h = int(total_seconds // 3600)
        m = int((total_seconds % 3600) // 60)
        return f"{h:02d}:{m:02d}"
    if isinstance(t, time):
        return t.strftime('%H:%M')
    return str(t)

def calc_end(row):
    """終了時間計算（元コードをそのまま）"""
    start = pd.to_datetime(row.get("start_dtime"), errors='coerce')
    man_hour = 0
    rest = 0

    mh = row.get("man_hour_time")
    if pd.notna(mh):
        if isinstance(mh, timedelta):
            man_hour = mh.total_seconds() / 3600
        elif isinstance(mh, (float, int)):
            man_hour = float(mh)
        elif isinstance(mh, time):
            man_hour = mh.hour + mh.minute / 60
        else:
            try:
                man_hour = float(str(mh))
            except:
                man_hour = 0

    rt = row.get("rest_time")
    if pd.notna(rt):
        if isinstance(rt, timedelta):
            rest = rt.total_seconds() / 60
        elif isinstance(rt, str) and ":" in rt:
            h, m = map(int, rt.split(":"))
            rest = h * 60 + m
        else:
            try:
                rest = float(rt) * 60
            except:
                rest = 0

    if man_hour > 0 and start is not pd.NaT:
        return start + timedelta(minutes=man_hour*60 - rest)
    return pd.to_datetime(row.get("end_dtime"), errors='coerce')


def export_worktime_personal(zexis_keiyaku_no, tanto_cd: str):
    """個人別明細をExcel出力"""
    engine = get_engine()

    # データ取得
    df_work = pd.read_sql(
        "SELECT * FROM weprj.work_record WHERE del_flg = false AND member_no = %s",
        con=engine,
        params=(tanto_cd,)
    )
    if df_work.empty:
        print(f"{tanto_cd} の勤務データがありません。")
        return

    df_man = pd.read_sql(
        "SELECT member_no, work_date, pj_cd, man_hour_time FROM weprj.man_hours WHERE del_flg = false AND member_no = %s",
        con=engine,
        params=(tanto_cd,)
    )

    df_member = pd.read_sql(
        """
        SELECT h.zexis_keiyaku_no, h.keiyaku_n, h.keiyaku_start_bi, h.keiyaku_end_bi,
               h.torihiki_rn, m.tanto_cd, m.tanto_n
        FROM d_keiyaku_seikyu_m m
        JOIN d_keiyaku_seikyu_h h ON m.zexis_keiyaku_no = h.zexis_keiyaku_no
        WHERE m.tanto_cd = %s
        """,
        con=engine,
        params=(tanto_cd,)
    )
    if df_member.empty:
        print(f"{tanto_cd} の契約情報がありません。")
        return

    df_rcv = pd.read_sql(
        "SELECT * FROM weprj.rcv_params WHERE order_no = %s AND del_flg = FALSE",
        con=engine,
        params=(zexis_keiyaku_no,)
    )
    if df_rcv.empty:
        print(f"{zexis_keiyaku_no} の受信パラメータがありません。")
        return

    df_holiday = pd.read_sql(
        "SELECT holiday_ymd FROM m_holiday WHERE del_flg = '0'",
        con=engine
    )
    holiday_set = set(pd.to_datetime(df_holiday['holiday_ymd']).dt.strftime("%Y%m%d"))

    # データ統合（drop_duplicates は削除）
    df = pd.merge(df_work, df_member, left_on='member_no', right_on='tanto_cd', how='left')
    df = pd.merge(df, df_man, on=["member_no", "work_date"], how="left")

    # Excel準備
    base_dir = Path(__file__).resolve().parent.parent
    template_path = base_dir / "templates" / "個人別明細A.xlsx"

    wb = load_workbook(template_path)
    ws = wb.active
    
    today = datetime.today()
    ws['E10'] = today.year
    ws['G10'] = today.month
    ws['I10'] = today.day

    data = df.iloc[0]
    rcv_data = df_rcv.iloc[0]

    start_date = pd.to_datetime(data['keiyaku_start_bi'])
    end_date = pd.to_datetime(data['keiyaku_end_bi'])

    ws['D1'] = data['torihiki_rn']
    ws['E7'] = data['zexis_keiyaku_no']
    ws['E8'] = data['keiyaku_n']
    ws['E9'] = f"{start_date:%Y/%m/%d}～{end_date:%Y/%m/%d}"
    ws['D14'] = data['tanto_n']
    ws['D2'] = ":" + rcv_data['ae_name']
    ws['R7'] = "業務担当者：" + rcv_data['pic_name']

    # 日付・時刻整形
    df['month'] = pd.to_datetime(df['work_date']).dt.month
    df['day'] = pd.to_datetime(df['work_date']).dt.day
    df['start_time'] = pd.to_datetime(df['start_dtime'], errors='coerce').dt.strftime('%H:%M')
    df['end_dtime'] = df.apply(calc_end, axis=1)
    df['end_time'] = pd.to_datetime(df['end_dtime'], errors='coerce').dt.strftime('%H:%M')

    # Excel書き込み
    start_row = 17
    df_to_write = df.head(31)

    for i, (_, row) in enumerate(df_to_write.iterrows()):
        row_idx = start_row + i
        ws.cell(row=row_idx, column=1, value=str(row['month']) if i == 0 else "")
        ws.cell(row=row_idx, column=2, value=row['day'])
        ws.cell(row=row_idx, column=3, value=row.get('weekday', ''))

        date_obj = pd.to_datetime(row['work_date'])
        weekday = date_obj.weekday()
        is_weekend = weekday >= 5
        date_str = date_obj.strftime("%Y%m%d")
        is_holiday = date_str in holiday_set

        status_cell = ws.cell(row=row_idx, column=4)
        u_col = ws.cell(row=row_idx, column=21)

        # D列：稼働 or 休暇
        start_dt = pd.to_datetime(row.get("start_dtime"), errors='coerce')
        if pd.isna(start_dt) or row['start_time'] == '':
            status_cell.value = "休暇"
            status_cell.font = Font(color="FF0000")
        else:
            status_cell.value = "稼働"
            status_cell.font = Font(color="0000FF")
            ws.cell(row=row_idx, column=5, value=row.get('start_time'))
            ws.cell(row=row_idx, column=6, value=row.get('end_time'))
            ws.cell(row=row_idx, column=7, value=time_to_str(row.get('rest_time')))
            ws.cell(row=row_idx, column=8, value=time_to_str(row.get('work_time')))
            ws.cell(row=row_idx, column=9, value=time_to_str(row.get('standard_time')))
            ws.cell(row=row_idx, column=10, value=time_to_str(row.get('over_time')))
            ws.cell(row=row_idx, column=11, value=time_to_str(row.get('night_time')))
            ws.cell(row=row_idx, column=12, value=time_to_str(row.get('paid_leave_time')))
            man_hour_val = row.get('man_hour_time')
            if pd.isna(man_hour_val) or man_hour_val == 0:
                ws.cell(row=row_idx, column=22, value="要確認！課金工数がありません。")
                ws.cell(row=row_idx, column=22).font = Font(color="FF0000")

        # U列：祝日/休み判定
        if is_holiday:
            u_col.value = "祝日"
        elif not is_weekend and pd.isna(start_dt):
            u_col.value = "休み"
        else:
            u_col.value = ""

    # 保存
    output_path = base_dir / "output" / f"個人別明細_{tanto_cd}.xlsx"
    wb.save(output_path)
    print(f"個人別明細出力完了: {output_path}")
