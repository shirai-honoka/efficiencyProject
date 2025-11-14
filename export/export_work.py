import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path
from db import get_engine


def export_work(zexis_keiyaku_no, df_staff, staff_list=None, pattern_code="PAT001"):
    engine = get_engine()

    # SQL（パラメータ化）
    sql_order = """
        SELECT 
            zexis_keiyaku_no,
            keiyaku_n,
            keiyaku_start_bi,
            keiyaku_end_bi,
            torihiki_rn
        FROM d_keiyaku_seikyu_h
        WHERE zexis_keiyaku_no = %(zexis)s
    """

    sql_rcv = """
        SELECT *
        FROM weprj.rcv_params
        WHERE order_no = %(zexis)s AND del_flg = FALSE
    """

    df_order = pd.read_sql(sql_order, con=engine, params={"zexis": zexis_keiyaku_no})
    df_rcv = pd.read_sql(sql_rcv, con=engine, params={"zexis": zexis_keiyaku_no})

    # データチェック
    if df_order.empty:
        print(f"[ERROR] 契約データなし: {zexis_keiyaku_no}")
        return
    if df_rcv.empty:
        print(f"[ERROR] rcv_params データなし: {zexis_keiyaku_no}")
        return

    data = df_order.iloc[0]
    rcv_data = df_rcv.iloc[0]

    # テンプレパス
    base_dir = Path(__file__).resolve().parent.parent
    template_path = base_dir / "templates" / "作業報告書1.xlsx"

    wb = load_workbook(template_path)
    ws = wb.active

    # 今日の日付
    today = datetime.today()
    ws.merge_cells("E10:F10")
    ws["E10"] = today.year
    ws["H10"] = today.month
    ws["J10"] = today.day

    # 契約期間（型安全）
    start_raw = data["keiyaku_start_bi"]
    end_raw = data["keiyaku_end_bi"]

    start_date = start_raw if isinstance(start_raw, str) else start_raw.strftime("%Y/%m/%d")
    end_date = end_raw if isinstance(end_raw, str) else end_raw.strftime("%Y/%m/%d")

    # header 書き込み
    ws["E9"] = f"{start_date}～{end_date}"
    ws["D1"] = data["torihiki_rn"]
    ws["E7"] = data["zexis_keiyaku_no"]
    ws["E8"] = data["keiyaku_n"]
    ws["D2"] = ":" + rcv_data["ae_name"]
    ws["Y7"] = "業務担当者：" + rcv_data["pic_name"]

    # スタッフ対象抽出
    df_target = df_staff[df_staff["order_no"] == zexis_keiyaku_no]

    if staff_list is not None:
        if isinstance(staff_list, str):
            staff_list = [staff_list]
        df_target = df_target[df_target["member_no"].isin(staff_list)]

    if df_target.empty:
        print(f"[WARN] スタッフデータなし: order={zexis_keiyaku_no}, staff={staff_list}")
        # 空のまま出力する（運用次第）
    
    # Excel の書き込み領域
    start_row = 15
    max_rows = 28

    # 初期化
    for i in range(start_row, start_row + max_rows):
        ws[f"B{i}"] = ""
        ws[f"G{i}"] = ""

    # 書き込み
    for idx, row in enumerate(df_target.itertuples(), start=start_row):
        if idx >= start_row + max_rows:
            break  # 行超え対策
        ws[f"B{idx}"] = getattr(row, "member_no", "")
        ws[f"G{idx}"] = getattr(row, "member_name", "")

    # 出力パス
    output_path = base_dir / "output" / f"作業報告書_{zexis_keiyaku_no}.xlsx"
    wb.save(output_path)

    print(f"作業報告書を出力: {output_path}")
