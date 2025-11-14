import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path
from db import get_engine


def export_kenshu(zexis_keiyaku_no: str, pattern_code: str = "PAT001"):
    """
    検収依頼書を出力する関数
    """
    engine = get_engine()

    sql = """
        SELECT
            zexis_keiyaku_no,
            keiyaku_n,
            keiyaku_start_bi,
            keiyaku_end_bi,
            torihiki_rn
        FROM d_keiyaku_seikyu_h
        WHERE zexis_keiyaku_no = %(zexis)s
    """

    sql_rcv = """
        SELECT *
        FROM weprj.rcv_params
        WHERE order_no = %(zexis)s AND del_flg = FALSE
    """

    df = pd.read_sql(sql, con=engine, params={"zexis": zexis_keiyaku_no})
    df_rcv = pd.read_sql(sql_rcv, con=engine, params={"zexis": zexis_keiyaku_no})

    if df.empty:
        print(f"[ERROR] 契約データなし: {zexis_keiyaku_no}")
        return
    if df_rcv.empty:
        print(f"[ERROR] rcv_params データなし: {zexis_keiyaku_no}")
        return

    data = df.iloc[0]
    rcv_data = df_rcv.iloc[0]

    # テンプレパス（絶対パスで安全化）
    base_dir = Path(__file__).resolve().parent.parent
    template_path = base_dir / "templates" / "検収依頼書1.xlsx"

    wb = load_workbook(template_path)
    ws = wb.active

    # 今日の日付
    today = datetime.today()
    ws.merge_cells("E10:F10")
    ws["E10"] = today.year
    ws["H10"] = today.month
    ws["J10"] = today.day

    # 契約期間（型安全化）
    start_raw = data["keiyaku_start_bi"]
    end_raw = data["keiyaku_end_bi"]

    if isinstance(start_raw, str):
        start_date = start_raw
    else:
        start_date = start_raw.strftime("%Y/%m/%d")

    if isinstance(end_raw, str):
        end_date = end_raw
    else:
        end_date = end_raw.strftime("%Y/%m/%d")

    ws["E9"] = f"{start_date}～{end_date}"
    ws["D1"] = data["torihiki_rn"]
    ws["B15"] = data["zexis_keiyaku_no"]
    ws["E15"] = data["keiyaku_n"]
    ws["D2"] = ":" + rcv_data["ae_name"]
    ws["Y7"] = "業務担当者：" + rcv_data["pic_name"]

    output_file = base_dir / "output" / f"検収依頼書_{zexis_keiyaku_no}.xlsx"
    wb.save(output_file)

    print(f"検収依頼書を出力: {output_file}")
